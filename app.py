import sys,re
from pathlib import Path
from datetime import datetime,timedelta
import fitz
from openpyxl import Workbook
from PySide6.QtCore import Qt,Signal
from PySide6.QtGui import QFont
from PySide6.QtWidgets import QApplication,QMainWindow,QWidget,QVBoxLayout,QHBoxLayout,QLabel,QPushButton,QTableWidget,QTableWidgetItem,QFileDialog,QMessageBox,QFrame,QAbstractItemView,QProgressBar

HEADERS=["Datei","RG-Nummer","Datum","Fälligkeit","Betrag Soll","Betrag Haben","Saldo","S/H Saldo","Status"]

def first(text,patterns):
    for p in patterns:
        m=re.search(p,text,re.I|re.M)
        if m:return m.group(1).strip()
    return ""

def amount(s):
    if not s:return None
    s=s.replace("€","").replace("EUR","").replace(" ","").strip()
    if "," in s:s=s.replace(".","").replace(",",".")
    try:return float(s)
    except:return None

def plus30(s):
    for f in ("%d.%m.%Y","%d.%m.%y"):
        try:return (datetime.strptime(s.replace("/","."),f)+timedelta(days=30)).strftime("%d.%m.%Y")
        except:pass
    return ""

def read_pdf(path):
    try:
        doc=fitz.open(path); text="\n".join(p.get_text("text") for p in doc); doc.close()
        rg=first(text,[r"(?:Rechnungsnummer|Rechnung(?:s)?-?Nr\\.?|RG-?Nr\\.?)\\s*[:#]?\\s*([A-Z0-9][A-Z0-9./_-]{2,})",r"\\bRG[- ]?([0-9]{3,})\\b"])
        date=first(text,[r"(?:Rechnungsdatum|Datum)\\s*[:.]?\\s*(\\d{1,2}[./-]\\d{1,2}[./-]\\d{2,4})"])
        due=first(text,[r"(?:Fälligkeit|fällig am|Zahlbar bis)\\s*[:.]?\\s*(\\d{1,2}[./-]\\d{1,2}[./-]\\d{2,4})"])
        soll=first(text,[r"(?:Betrag Soll|Sollbetrag|Rechnungsbetrag|Gesamtbetrag|Zu zahlen|Endbetrag)\\s*[:.]?\\s*([0-9][0-9.]*,\\d{2}|[0-9]+[.,]\\d{2})"])
        haben=first(text,[r"(?:Betrag Haben|Haben)\\s*[:.]?\\s*([0-9][0-9.]*,\\d{2}|[0-9]+[.,]\\d{2})"])
        saldo=first(text,[r"(?:Saldo)\\s*[:.]?\\s*(-?[0-9][0-9.]*,\\d{2}|-?[0-9]+[.,]\\d{2})"])
        sh=first(text,[r"(?:S/H Saldo|S/H-Saldo)\\s*[:.]?\\s*(-?[0-9][0-9.]*,\\d{2}|-?[0-9]+[.,]\\d{2})"])
        if date and not due: due=plus30(date)
        return dict(file=Path(path).name,rg=rg,date=date,due=due,soll=amount(soll),haben=amount(haben),saldo=amount(saldo),sh=amount(sh),status="✓ Erkannt" if text.strip() else "⚠ Kein Text",path=path)
    except Exception as e:
        return dict(file=Path(path).name,rg="",date="",due="",soll=None,haben=None,saldo=None,sh=None,status="⚠ "+str(e),path=path)

class Drop(QFrame):
    dropped=Signal(list)
    def __init__(self):
        super().__init__();self.setAcceptDrops(True);self.setObjectName("drop")
        l=QVBoxLayout(self);a=QLabel("PDFs hier hineinziehen");a.setAlignment(Qt.AlignCenter);a.setObjectName("dropTitle")
        b=QLabel("oder „PDF hinzufügen“ wählen");b.setAlignment(Qt.AlignCenter);b.setObjectName("dropSub")
        l.addStretch();l.addWidget(a);l.addWidget(b);l.addStretch()
    def dragEnterEvent(self,e):
        if any(u.toLocalFile().lower().endswith(".pdf") for u in e.mimeData().urls()):e.acceptProposedAction()
    def dropEvent(self,e):
        p=[u.toLocalFile() for u in e.mimeData().urls() if u.toLocalFile().lower().endswith(".pdf")]
        if p:self.dropped.emit(p)

class Main(QMainWindow):
    def __init__(self):
        super().__init__();self.rows=[];self.setWindowTitle("PDF → Excel");self.resize(1200,760);self.ui()
    def ui(self):
        c=QWidget();self.setCentralWidget(c);l=QVBoxLayout(c);l.setContentsMargins(28,24,28,24);l.setSpacing(14)
        top=QHBoxLayout();t=QLabel("PDF → Excel");t.setObjectName("title");self.count=QLabel("0 PDFs");top.addWidget(t);top.addStretch();top.addWidget(self.count);l.addLayout(top)
        self.drop=Drop();self.drop.dropped.connect(self.add);l.addWidget(self.drop)
        bar=QHBoxLayout()
        for text,fn in [("+ PDF hinzufügen",self.pick),("Ausgewählte löschen",self.remove),("Alles entfernen",self.clear)]:
            b=QPushButton(text);b.clicked.connect(fn);bar.addWidget(b)
        bar.addStretch();b=QPushButton("Nach Excel exportieren");b.setObjectName("export");b.clicked.connect(self.export);bar.addWidget(b);l.addLayout(bar)
        self.progress=QProgressBar();self.progress.setVisible(False);l.addWidget(self.progress)
        self.table=QTableWidget(0,len(HEADERS));self.table.setHorizontalHeaderLabels(HEADERS);self.table.setSelectionBehavior(QAbstractItemView.SelectRows);self.table.setAlternatingRowColors(True);self.table.horizontalHeader().setStretchLastSection(True);self.table.setColumnWidth(0,240);l.addWidget(self.table,1)
        self.info=QLabel("Noch keine PDFs hinzugefügt.");l.addWidget(self.info)
    def pick(self):
        p,_=QFileDialog.getOpenFileNames(self,"PDFs auswählen","","PDF-Dateien (*.pdf)")
        if p:self.add(p)
    def add(self,paths):
        old={r["file"] for r in self.rows};paths=[p for p in paths if Path(p).name not in old]
        self.progress.setVisible(True);self.progress.setRange(0,len(paths))
        for i,p in enumerate(paths,1):self.rows.append(read_pdf(p));self.progress.setValue(i);QApplication.processEvents()
        self.progress.setVisible(False);self.refresh()
    def refresh(self):
        self.table.setRowCount(len(self.rows))
        for r,d in enumerate(self.rows):
            vals=[d["file"],d["rg"],d["date"],d["due"],self.money(d["soll"]),self.money(d["haben"]),self.money(d["saldo"]),self.money(d["sh"]),d["status"]]
            for c,v in enumerate(vals):
                x=QTableWidgetItem(str(v));x.setFlags(x.flags()|Qt.ItemIsEditable);self.table.setItem(r,c,x)
        n=len(self.rows);good=sum("✓" in x["status"] for x in self.rows);self.count.setText(f"{n} PDF{'s' if n!=1 else ''}");self.info.setText(f"{n} PDFs · {good} erkannt · {n-good} mit Hinweis")
    def remove(self):
        for i in sorted([x.row() for x in self.table.selectionModel().selectedRows()],reverse=True):del self.rows[i]
        self.refresh()
    def clear(self):self.rows=[];self.refresh()
    @staticmethod
    def money(v):
        return "" if v is None else f"{v:,.2f} €".replace(",","X").replace(".",",").replace("X",".")
    def export(self):
        if not self.rows:return QMessageBox.information(self,"Keine Daten","Bitte zuerst PDFs hinzufügen.")
        for r,d in enumerate(self.rows):
            d["file"]=self.table.item(r,0).text();d["rg"]=self.table.item(r,1).text();d["date"]=self.table.item(r,2).text();d["due"]=self.table.item(r,3).text()
        p,_=QFileDialog.getSaveFileName(self,"Excel-Datei speichern","PDF_Export.xlsx","Excel-Dateien (*.xlsx)")
        if not p:return
        wb=Workbook();ws=wb.active;ws.title="PDF Export";ws.append(HEADERS)
        for d in self.rows:ws.append([d["file"],d["rg"],d["date"],d["due"],d["soll"],d["haben"],d["saldo"],d["sh"],d["status"]])
        for x in ws[1]:x.font=x.font.copy(bold=True)
        for col in "EFGH":
            for x in ws[col][1:]:x.number_format='#,##0.00 "€"'
        ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
        for i,w in enumerate([30,18,14,14,18,18,16,16,18],1):ws.column_dimensions[chr(64+i)].width=w
        wb.save(p);QMessageBox.information(self,"Export fertig","Excel-Datei wurde gespeichert.")

STYLE="""QMainWindow,QWidget{background:#151515;color:#eee}#title{font-size:28px;font-weight:700}#drop{background:#1d1d1d;border:1px dashed #555;border-radius:14px;min-height:135px}#dropTitle{font-size:20px;font-weight:600}#dropSub{color:#888}QPushButton{background:#262626;color:#eee;border:1px solid #3b3b3b;border-radius:8px;padding:10px 16px}QPushButton:hover{background:#303030}#export{background:#eee;color:#111;font-weight:700}QTableWidget{background:#191919;alternate-background-color:#1e1e1e;border:1px solid #303030;gridline-color:#292929;selection-background-color:#343434}QHeaderView::section{background:#222;color:#aaa;padding:9px;font-weight:600}QProgressBar{border:none;background:#252525;height:4px}QProgressBar::chunk{background:#aaa}"""

if __name__=="__main__":
    app=QApplication(sys.argv);app.setStyleSheet(STYLE);app.setFont(QFont("Segoe UI",10));w=Main();w.show();sys.exit(app.exec())
